"""Unit tests for business logic that needs no HTTP layer."""
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import jwt
import pytest
from fastapi import HTTPException
from pydantic import ValidationError

from app.core.security import (
    as_utc, create_access_token, create_refresh_token, decode_token,
    hash_opaque_token, hash_password, verify_password,
)
from app.models import (
    Cart, CartItem, Coupon, Inventory, Product, ShippingMethod, TaxRate,
)
from app.schemas.auth import UserRegister, validate_phone
from app.schemas.order import CheckoutRequest
from app.services import exports
from app.services.pricing import calculate_tax, money, price_cart


# ----------------------------------------------------------- validation ----

@pytest.mark.parametrize("raw,expected", [
    ("+919876543210", "+919876543210"),
    ("+91 98765-43210", "+919876543210"),
    (None, None),
    ("  ", None),
])
def test_phone_normalisation(raw, expected):
    assert validate_phone(raw) == expected


@pytest.mark.parametrize("raw", ["9876543210", "+0123456789", "+12", "abc", "+91987654321012345"])
def test_phone_rejects_invalid_numbers(raw):
    with pytest.raises(ValueError):
        validate_phone(raw)


def test_registration_schema_rules():
    ok = UserRegister(full_name="Ann Lee", email="ann@example.com", password="Passw0rd!")
    assert ok.phone is None
    for bad in ({"password": "short1"}, {"password": "lettersonly"}, {"password": "12345678"}, {"full_name": "A"}):
        with pytest.raises(ValidationError):
            UserRegister(**{"full_name": "Ann Lee", "email": "ann@example.com", "password": "Passw0rd!", **bad})


def test_checkout_request_needs_address_id_or_inline_address():
    with pytest.raises(ValidationError):
        CheckoutRequest(shipping_method_id=1)
    assert CheckoutRequest(shipping_method_id=1, shipping_address_id=5).shipping_address_id == 5
    inline = CheckoutRequest(
        shipping_method_id=1, shipping_full_name="Ann Lee", shipping_phone="+919876543210",
        shipping_address_line1="1 Road", shipping_city="Pune", shipping_state="MH", shipping_postal_code="411001",
    )
    assert inline.shipping_country == "IN"


# ------------------------------------------------------------- security ----

def test_password_hash_roundtrip():
    hashed = hash_password("Secret#123")
    assert hashed != "Secret#123" and verify_password("Secret#123", hashed)
    assert not verify_password("secret#123", hashed)


def test_tokens_carry_type_and_session():
    access = decode_token(create_access_token(7, "sess1"), "access")
    assert access["sub"] == "7" and access["sid"] == "sess1"
    refresh_token, jti, expires = create_refresh_token(7, "sess1")
    assert decode_token(refresh_token, "refresh")["jti"] == jti
    assert expires > datetime.now(timezone.utc)
    with pytest.raises(jwt.InvalidTokenError):
        decode_token(refresh_token, "access")


def test_expired_token_is_rejected(monkeypatch):
    from app.core.config import settings
    monkeypatch.setattr(settings, "access_token_expire_minutes", -1)
    with pytest.raises(jwt.ExpiredSignatureError):
        decode_token(create_access_token(1, "s"), "access")


def test_opaque_token_hash_is_stable_and_not_reversible():
    assert hash_opaque_token("abc") == hash_opaque_token("abc") != "abc"
    assert len(hash_opaque_token("abc")) == 64


def test_as_utc_handles_naive_and_aware_values():
    naive = datetime(2026, 1, 1, 12, 0)
    assert as_utc(naive).tzinfo is not None
    aware = datetime(2026, 1, 1, 12, 0, tzinfo=timezone(timedelta(hours=5)))
    assert as_utc(aware).hour == 7


# -------------------------------------------------------------- exports ----

SAMPLE = {"report": "t", "columns": ["Name", "Qty"], "rows": [["=HYPERLINK(\"x\")", 1], ["ok", 2]], "summary": {"rows": 2}}


def test_csv_export_neutralises_formulas():
    text = exports.to_csv(SAMPLE).decode()
    assert text.splitlines()[0] == "Name,Qty"
    assert "'=HYPERLINK" in text


def test_xlsx_and_pdf_exports_are_valid_files():
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(exports.to_xlsx(SAMPLE)))
    assert workbook.sheetnames == ["t", "Summary"]
    assert workbook["t"]["A2"].value.startswith("'=")
    assert exports.to_pdf(SAMPLE)[:4] == b"%PDF"


def test_unsupported_export_format():
    with pytest.raises(ValueError):
        exports.render(SAMPLE, "docx")


# -------------------------------------------------------------- pricing ----

def _product(db, sku, price, category_id=None):
    product = Product(sku=sku, name=sku, category="c", category_id=category_id, price=Decimal(price))
    db.add(product)
    db.flush()
    db.add(Inventory(product_id=product.id, quantity_on_hand=100, low_stock_threshold=1))
    return product


def test_money_rounds_half_up():
    assert money("1.005") == Decimal("1.01") and money(2) == Decimal("2.00")


def test_tax_exclusive_and_inclusive(db):
    product = _product(db, "A", "100")
    db.add(TaxRate(name="Ex", code="EX", rate_percent=Decimal("18"), country_code="IN", is_inclusive=False))
    db.commit()
    assert calculate_tax(db, Decimal("100.00"), product, "IN", None) == (Decimal("100.00"), Decimal("18.00"))
    assert calculate_tax(db, Decimal("100.00"), product, "US", None) == (Decimal("100.00"), Decimal("0.00"))

    db.query(TaxRate).delete()
    db.add(TaxRate(name="Inc", code="INC", rate_percent=Decimal("18"), country_code="IN", is_inclusive=True))
    db.commit()
    base, tax = calculate_tax(db, Decimal("118.00"), product, "IN", None)
    assert (base, tax) == (Decimal("100.00"), Decimal("18.00"))


def test_category_specific_tax_beats_default(db):
    from app.models import Category
    category = Category(name="Food", slug="food")
    db.add(category)
    db.flush()
    product = _product(db, "F", "100", category.id)
    db.add_all([
        TaxRate(name="Default", code="D", rate_percent=Decimal("18"), country_code="IN"),
        TaxRate(name="Food", code="F5", rate_percent=Decimal("5"), country_code="IN", category_id=category.id),
    ])
    db.commit()
    assert calculate_tax(db, Decimal("100.00"), product, "IN", None)[1] == Decimal("5.00")


def _cart_with(db, user, products, quantity=1):
    cart = Cart(user_id=user.id)
    db.add(cart)
    db.flush()
    items = [CartItem(cart_id=cart.id, product_id=p.id, quantity=quantity) for p in products]
    db.add_all(items)
    db.commit()
    return items


def test_inclusive_tax_is_not_double_counted(db, users):
    product = _product(db, "I", "118")
    db.add(TaxRate(name="Inc", code="INC", rate_percent=Decimal("18"), country_code="IN", is_inclusive=True))
    db.commit()
    result = price_cart(db, _cart_with(db, users[next(iter(users))], [product]))
    assert result.grand_total == Decimal("118.00") and result.tax_total == Decimal("18.00")


def test_fixed_coupon_is_split_exactly_across_lines(db, users):
    products = [_product(db, f"S{i}", "10.00") for i in range(3)]
    now = datetime.now(timezone.utc)
    db.add(Coupon(code="FIX10", discount_type="FIXED", discount_value=Decimal("10.00"),
                  min_order_amount=Decimal("0.00"), starts_at=now - timedelta(days=1), expires_at=now + timedelta(days=1)))
    db.commit()
    result = price_cart(db, _cart_with(db, users[next(iter(users))], products), coupon_code="fix10")
    assert result.coupon_discount == Decimal("10.00")
    assert sum(line.coupon_share for line in result.lines) == Decimal("10.00")
    assert result.grand_total == Decimal("20.00")


def test_coupon_cannot_exceed_order_value(db, users):
    product = _product(db, "C", "5.00")
    now = datetime.now(timezone.utc)
    db.add(Coupon(code="BIG", discount_type="FIXED", discount_value=Decimal("50.00"),
                  min_order_amount=Decimal("0.00"), starts_at=now - timedelta(days=1), expires_at=now + timedelta(days=1)))
    db.commit()
    result = price_cart(db, _cart_with(db, users[next(iter(users))], [product]), coupon_code="BIG")
    assert result.grand_total == Decimal("0.00") and result.coupon_discount == Decimal("5.00")


@pytest.mark.parametrize("kwargs,detail", [
    ({"minimum": "100.00"}, "minimum subtotal"),
    ({"expired": True}, "validity"),
    ({"used_up": True}, "usage limit"),
    ({"inactive": True}, "unavailable"),
])
def test_coupon_rules(db, users, kwargs, detail):
    product = _product(db, "R", "10.00")
    now = datetime.now(timezone.utc)
    db.add(Coupon(
        code="RULE", discount_type="FIXED", discount_value=Decimal("1.00"),
        min_order_amount=Decimal(kwargs.get("minimum", "0.00")),
        starts_at=now - timedelta(days=2),
        expires_at=now - timedelta(days=1) if kwargs.get("expired") else now + timedelta(days=1),
        usage_limit=1 if kwargs.get("used_up") else None,
        used_count=1 if kwargs.get("used_up") else 0,
        is_active=not kwargs.get("inactive"),
    ))
    db.commit()
    with pytest.raises(HTTPException) as error:
        price_cart(db, _cart_with(db, users[next(iter(users))], [product]), coupon_code="RULE")
    assert error.value.status_code == 400 and detail in error.value.detail


def test_free_shipping_threshold(db, users):
    product = _product(db, "FS", "60.00")
    method = ShippingMethod(code="T", name="T", base_fee=Decimal("40.00"), free_shipping_threshold=Decimal("100.00"))
    db.add(method)
    db.commit()
    items = _cart_with(db, users[next(iter(users))], [product], quantity=1)
    assert price_cart(db, items, shipping_method_id=method.id).shipping_total == Decimal("40.00")
    items[0].quantity = 2
    db.commit()
    assert price_cart(db, items, shipping_method_id=method.id).shipping_total == Decimal("0.00")


def test_inactive_product_or_empty_cart_is_rejected(db, users):
    product = _product(db, "X", "10.00")
    items = _cart_with(db, users[next(iter(users))], [product])
    product.is_active = False
    db.commit()
    with pytest.raises(HTTPException) as error:
        price_cart(db, items)
    assert error.value.status_code == 409
    with pytest.raises(HTTPException):
        price_cart(db, [])


def test_insufficient_stock_is_rejected(db, users):
    product = _product(db, "LOW", "10.00")
    items = _cart_with(db, users[next(iter(users))], [product], quantity=101)
    with pytest.raises(HTTPException) as error:
        price_cart(db, items)
    assert error.value.status_code == 409
    assert price_cart(db, items, check_stock=False).item_count == 101
