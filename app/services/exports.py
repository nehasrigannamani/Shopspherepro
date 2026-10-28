"""Render a report result as CSV, Excel or PDF bytes."""
import csv
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMATS = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}
PDF_ROW_LIMIT = 1000


def _safe_cell(value):
    """Neutralise spreadsheet formula injection in text cells."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def to_csv(result: dict) -> bytes:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(result["columns"])
    for row in result["rows"]:
        writer.writerow([_safe_cell(c) for c in row])
    return out.getvalue().encode("utf-8")


def to_xlsx(result: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = result["report"][:31]
    ws.append(result["columns"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in result["rows"]:
        ws.append([_safe_cell(c) for c in row])
    for idx, column in enumerate(result["columns"], start=1):
        width = max(
            [len(str(column))] + [len(str(r[idx - 1])) for r in result["rows"][:200]]
        )
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(width + 2, 50)

    summary = wb.create_sheet("Summary")
    for key, value in result["summary"].items():
        summary.append([str(key), str(value)])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def to_pdf(result: dict) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), title=f"{result['report']} report")
    styles = getSampleStyleSheet()
    story = [
        Paragraph(f"ShopSphere Pro - {result['report']} report", styles["Title"]),
        Paragraph(
            "Generated " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"),
            styles["Normal"],
        ),
        Spacer(1, 8),
    ]
    for key, value in result["summary"].items():
        story.append(Paragraph(f"<b>{key}</b>: {value}", styles["Normal"]))
    story.append(Spacer(1, 10))

    rows = result["rows"][:PDF_ROW_LIMIT]
    data = [result["columns"]] + [[str(c) for c in r] for r in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    if len(result["rows"]) > PDF_ROW_LIMIT:
        story.append(Spacer(1, 6))
        story.append(
            Paragraph(
                f"Showing first {PDF_ROW_LIMIT} of {len(result['rows'])} rows; "
                "use CSV or Excel for the full data.",
                styles["Italic"],
            )
        )
    doc.build(story)
    return out.getvalue()


def render(result: dict, fmt: str) -> bytes:
    if fmt == "csv":
        return to_csv(result)
    if fmt == "xlsx":
        return to_xlsx(result)
    if fmt == "pdf":
        return to_pdf(result)
    raise ValueError(f"Unsupported export format: {fmt}")
